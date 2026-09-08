import sqlite3
import os
import sys
import shutil
import glob
from io import BytesIO
from collections import defaultdict

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

try:
    import jdatetime
    HAS_JDATETIME = True
except ImportError:
    HAS_JDATETIME = False

try:
    import matplotlib
    matplotlib.use('TkAgg')
    import matplotlib.pyplot as plt
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
    import matplotlib.font_manager as fm
    
    # تنظیم فونت فارسی برای matplotlib
    def setup_persian_font_for_matplotlib():
        """تنظیم فونت فارسی برای matplotlib"""
        # لیست فونت‌های فارسی که هم فارسی و هم لاتین را پشتیبانی می‌کنند
        # Vazirmatn و IRANSans بهترین گزینه‌ها هستند
        persian_fonts = [
            'Vazirmatn', 'IRANSans', 'Vazir', 'Shabnam', 'Sahel',
            'B Nazanin', 'B Mitra', 'B Yekan', 'Tahoma', 'Arial'
        ]
        
        # پیدا کردن فونت‌های نصب شده
        available_fonts = [f.name for f in fm.fontManager.ttflist]
        
        # اولویت با Vazirmatn است چون هم فارسی و هم لاتین را خوب پشتیبانی می‌کند
        preferred_fonts = ['Vazirmatn', 'IRANSans', 'Vazir', 'Shabnam', 'Sahel']
        
        for font in preferred_fonts:
            if font in available_fonts:
                return font
        
        # اگر هیچ فونت فارسی پیدا نشد، از فونت پیش‌فرض استفاده کن
        return 'DejaVu Sans'
    
    PERSIAN_FONT = setup_persian_font_for_matplotlib()
    
    # تنظیم فونت برای matplotlib
    matplotlib.rcParams['font.family'] = PERSIAN_FONT
    matplotlib.rcParams['axes.unicode_minus'] = False
    matplotlib.rcParams['font.size'] = 10
    
    # برای اطمینان از اینکه فونت به درستی تنظیم شده
    if PERSIAN_FONT != 'DejaVu Sans':
        # پیدا کردن مسیر فونت و تنظیم آن
        for f in fm.fontManager.ttflist:
            if f.name == PERSIAN_FONT:
                matplotlib.rcParams['font.family'] = [PERSIAN_FONT, 'DejaVu Sans', 'sans-serif']
                break
    
    HAS_MPL = True
except ImportError:
    HAS_MPL = False

try:
    import openpyxl
    from openpyxl.styles import Font as XLFont, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    HAS_RESHAPE = True
except ImportError:
    HAS_RESHAPE = False

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

BG_MAIN = "#0f1220"
BG_CARD = "#1a1e33"
BG_CARD_ALT = "#20264a"
FG_TEXT = "#e8e9f3"
FG_MUTED = "#9096b5"
ACCENT = "#7c5cff"
ACCENT2 = "#4fd1c5"
GREEN = "#3ddc97"
RED = "#ff6b6b"
BLUE = "#4d96ff"
YELLOW = "#ffd166"
FONT_FAMILY = "Segoe UI" if sys.platform.startswith("win") else "Arial"

JALALI_MONTH_NAMES = [
    "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
    "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند",
]
JALALI_MONTH_SHORT = [
    "فرو", "ارد", "خرد", "تیر", "مرد", "شهر",
    "مهر", "آبا", "آذر", "دی", "بهم", "اسف",
]

EXPENSE_CATEGORIES = [
    "خوراک و خواربار", "حمل‌ونقل", "مسکن و اجاره", "قبوض",
    "پوشاک", "سلامت و درمان", "سرگرمی", "آموزش", "متفرقه",
]
INCOME_CATEGORIES = [
    "حقوق", "فریلنسری", "سرمایه‌گذاری", "هدیه", "پاداش", "متفرقه",
]

DB_PATH = os.path.join(SCRIPT_DIR, "expenses.db")
BACKUP_DIR = os.path.join(SCRIPT_DIR, "backups")
MAX_BACKUPS = 20

def jtoday():
    if HAS_JDATETIME:
        return jdatetime.date.today()
    class _Fallback:
        year, month, day = 1403, 1, 1
    return _Fallback()


def jnow_str():
    if HAS_JDATETIME:
        n = jdatetime.datetime.now()
        return f"{n.year:04d}{n.month:02d}{n.day:02d}-{n.hour:02d}{n.minute:02d}{n.second:02d}"
    import time
    return time.strftime("%Y%m%d-%H%M%S")


def parse_jalali(text):
    if not HAS_JDATETIME:
        raise RuntimeError("کتابخانه jdatetime نصب نیست")
    parts = text.strip().split("-")
    if len(parts) != 3:
        raise ValueError("فرمت باید YYYY-MM-DD باشد")
    y, m, d = (int(p) for p in parts)
    return jdatetime.date(y, m, d)


def jalali_str(jdate):
    return f"{jdate.year:04d}-{jdate.month:02d}-{jdate.day:02d}"


def month_label(year, month):
    return f"{JALALI_MONTH_NAMES[month-1]} {year}"


def fmt_money(v):
    try:
        return f"{v:,.0f}"
    except Exception:
        return str(v)


def rtl(text):
    """برای رندر درست فارسی در matplotlib/reportlab: shaping + راست‌به‌چپ."""
    if not text:
        return text
    if HAS_RESHAPE:
        try:
            return get_display(arabic_reshaper.reshape(str(text)))
        except Exception:
            return text
    return text

_PERSIAN_FONT_NAME = None


def find_persian_font():
    """به‌دنبال یک فونت TTF فارسی می‌گردد: اول کنار اسکریپت، بعد مسیرهای رایج سیستم."""
    candidates = []
    candidates += glob.glob(os.path.join(SCRIPT_DIR, "*.ttf"))
    candidates += glob.glob(os.path.join(SCRIPT_DIR, "fonts", "*.ttf"))
    if sys.platform.startswith("win"):
        windir = os.environ.get("WINDIR", r"C:\Windows")
        candidates += [
            os.path.join(windir, "Fonts", "tahoma.ttf"),
            os.path.join(windir, "Fonts", "arial.ttf"),
        ]
    elif sys.platform == "darwin":
        candidates += [
            "/Library/Fonts/Tahoma.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
        ]
    else:
        candidates += [
            "/usr/share/fonts/truetype/vazir/Vazir.ttf",
            "/usr/share/fonts/truetype/vazirmatn/Vazirmatn-Regular.ttf",
        ]
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return None


def register_persian_font():
    global _PERSIAN_FONT_NAME
    if not HAS_REPORTLAB:
        return None
    if _PERSIAN_FONT_NAME:
        return _PERSIAN_FONT_NAME
    path = find_persian_font()
    if path:
        try:
            pdfmetrics.registerFont(TTFont("Persian", path))
            _PERSIAN_FONT_NAME = "Persian"
            return _PERSIAN_FONT_NAME
        except Exception:
            pass
    _PERSIAN_FONT_NAME = "Helvetica"
    return _PERSIAN_FONT_NAME

def prune_backups(keep=MAX_BACKUPS):
    files = sorted(glob.glob(os.path.join(BACKUP_DIR, "expenses_*.db")))
    excess = len(files) - keep
    for f in files[:max(0, excess)]:
        try:
            os.remove(f)
        except OSError:
            pass


def backup_now():
    if not os.path.isfile(DB_PATH):
        return None
    os.makedirs(BACKUP_DIR, exist_ok=True)
    dest = os.path.join(BACKUP_DIR, f"expenses_{jnow_str()}.db")
    shutil.copy2(DB_PATH, dest)
    prune_backups()
    return dest


def auto_backup_if_needed():
    """اگر امروز پشتیبانی گرفته نشده، یکی می‌گیرد."""
    os.makedirs(BACKUP_DIR, exist_ok=True)
    today_prefix = jnow_str().split("-")[0]
    existing = glob.glob(os.path.join(BACKUP_DIR, f"expenses_{today_prefix}-*.db"))
    if not existing:
        backup_now()


def list_backups():
    files = sorted(glob.glob(os.path.join(BACKUP_DIR, "expenses_*.db")), reverse=True)
    out = []
    for f in files:
        size_kb = os.path.getsize(f) / 1024
        out.append((f, os.path.basename(f), f"{size_kb:.0f} KB"))
    return out

class Database:
    def __init__(self, path=DB_PATH):
        self.path = path
        self._connect()

    def _connect(self):
        self.conn = sqlite3.connect(self.path)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                kind TEXT NOT NULL CHECK(kind IN ('income','expense')),
                amount REAL NOT NULL,
                category TEXT NOT NULL,
                note TEXT,
                tdate TEXT NOT NULL
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS budgets (
                category TEXT PRIMARY KEY,
                amount REAL NOT NULL
            )
        """)
        self.conn.commit()

    def reconnect(self):
        try:
            self.conn.close()
        except Exception:
            pass
        self._connect()

    def add(self, kind, amount, category, note, tdate):
        self.conn.execute(
            "INSERT INTO transactions (kind, amount, category, note, tdate) VALUES (?,?,?,?,?)",
            (kind, amount, category, note, tdate),
        )
        self.conn.commit()

    def update(self, tid, kind, amount, category, note, tdate):
        self.conn.execute(
            "UPDATE transactions SET kind=?, amount=?, category=?, note=?, tdate=? WHERE id=?",
            (kind, amount, category, note, tdate, tid),
        )
        self.conn.commit()

    def delete(self, tid):
        self.conn.execute("DELETE FROM transactions WHERE id=?", (tid,))
        self.conn.commit()

    def get(self, tid):
        cur = self.conn.execute(
            "SELECT id, kind, amount, category, note, tdate FROM transactions WHERE id=?", (tid,)
        )
        return cur.fetchone()

    def all(self):
        cur = self.conn.execute(
            "SELECT id, kind, amount, category, note, tdate FROM transactions ORDER BY tdate DESC, id DESC"
        )
        return cur.fetchall()

    def for_month(self, year, month):
        prefix = f"{year:04d}-{month:02d}"
        cur = self.conn.execute(
            "SELECT id, kind, amount, category, note, tdate FROM transactions "
            "WHERE tdate LIKE ? ORDER BY tdate DESC, id DESC",
            (prefix + "%",),
        )
        return cur.fetchall()

    def for_year(self, year):
        prefix = f"{year:04d}"
        cur = self.conn.execute(
            "SELECT id, kind, amount, category, note, tdate FROM transactions "
            "WHERE tdate LIKE ? ORDER BY tdate DESC, id DESC",
            (prefix + "%",),
        )
        return cur.fetchall()

    def years_present(self):
        cur = self.conn.execute("SELECT DISTINCT substr(tdate,1,4) FROM transactions ORDER BY 1")
        years = [int(r[0]) for r in cur.fetchall()]
        ty = jtoday().year
        if ty not in years:
            years.append(ty)
        return sorted(set(years))

    # --- بودجه ---
    def get_budgets(self):
        cur = self.conn.execute("SELECT category, amount FROM budgets")
        return {r[0]: r[1] for r in cur.fetchall()}

    def set_budget(self, category, amount):
        self.conn.execute(
            "INSERT INTO budgets (category, amount) VALUES (?,?) "
            "ON CONFLICT(category) DO UPDATE SET amount=excluded.amount",
            (category, amount),
        )
        self.conn.commit()

class StatCard(tk.Frame):
    def __init__(self, master, title, color, **kw):
        super().__init__(master, bg=BG_CARD, highlightbackground=color,
                          highlightthickness=1, bd=0, **kw)
        self.title = title
        self.color = color
        tk.Label(self, text=title, bg=BG_CARD, fg=FG_MUTED,
                  font=(FONT_FAMILY, 11)).pack(anchor="w", padx=16, pady=(14, 0))
        self.value_lbl = tk.Label(self, text="0", bg=BG_CARD, fg=color,
                                   font=(FONT_FAMILY, 22, "bold"))
        self.value_lbl.pack(anchor="w", padx=16, pady=(2, 14))

    def set_value(self, text):
        self.value_lbl.config(text=text)


class BudgetRow(tk.Frame):
    """یک ردیف بودجه برای یک دسته‌بندی هزینه."""

    def __init__(self, master, category, on_change, **kw):
        super().__init__(master, bg=BG_CARD, **kw)
        self.category = category
        self.amount_var = tk.StringVar(value="0")

        tk.Label(self, text=category, bg=BG_CARD, fg=FG_TEXT, width=18, anchor="e",
                  font=(FONT_FAMILY, 10, "bold")).grid(row=0, column=4, sticky="e", padx=8, pady=8)

        entry = ttk.Entry(self, textvariable=self.amount_var, width=12, justify="center")
        entry.grid(row=0, column=3, padx=8)
        entry.bind("<FocusOut>", lambda e: on_change())

        self.progress = ttk.Progressbar(self, length=220, maximum=100, mode="determinate")
        self.progress.grid(row=0, column=2, padx=8)

        self.status_lbl = tk.Label(self, text="", bg=BG_CARD, fg=FG_MUTED,
                                    font=(FONT_FAMILY, 9), width=28, anchor="w")
        self.status_lbl.grid(row=0, column=1, sticky="w", padx=8)

        self.columnconfigure(0, weight=1)

    def set_budget(self, amount):
        self.amount_var.set(fmt_money(amount) if amount else "0")

    def get_budget(self):
        raw = self.amount_var.get().strip().replace(",", "")
        try:
            return max(0.0, float(raw)) if raw else 0.0
        except ValueError:
            return 0.0

    def set_spent(self, spent, budget):
        if budget > 0:
            pct = (spent / budget) * 100
            self.progress["value"] = min(pct, 100)
            color = GREEN if pct < 80 else (YELLOW if pct < 100 else RED)
            self.status_lbl.config(
                text=f"{fmt_money(spent)} / {fmt_money(budget)}  ({pct:.0f}%)", fg=color)
        else:
            self.progress["value"] = 0
            self.status_lbl.config(text=f"{fmt_money(spent)} (بدون بودجه)", fg=FG_MUTED)

class ExpenseApp(tk.Tk):
    def __init__(self):
        super().__init__()
        auto_backup_if_needed()
        self.db = Database()
        self.title("مدیریت خرج و مخارج")
        self.geometry("1150x720")
        self.minsize(940, 620)
        self.configure(bg=BG_MAIN)

        if not HAS_JDATETIME:
            messagebox.showwarning(
                "کتابخانه‌ی jdatetime یافت نشد",
                "برای کارکرد کامل تاریخ شمسی نصب کنید:\npip install jdatetime")

        today = jtoday()
        self.sel_year = tk.IntVar(value=today.year)
        self.sel_month = tk.IntVar(value=today.month)
        self.budget_rows = {}
        self.edit_dialog = None

        self._setup_style()
        self._build_layout()
        self.refresh_all()

    def _setup_style(self):
        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure("TNotebook", background=BG_MAIN, borderwidth=0)
        style.configure("TNotebook.Tab", background=BG_CARD, foreground=FG_MUTED,
                         padding=(16, 9), font=(FONT_FAMILY, 10, "bold"))
        style.map("TNotebook.Tab",
                  background=[("selected", ACCENT)],
                  foreground=[("selected", "#ffffff")])

        style.configure("TFrame", background=BG_MAIN)
        style.configure("Card.TFrame", background=BG_CARD)
        style.configure("TLabel", background=BG_MAIN, foreground=FG_TEXT, font=(FONT_FAMILY, 10))
        style.configure("Muted.TLabel", background=BG_MAIN, foreground=FG_MUTED)
        style.configure("Card.TLabel", background=BG_CARD, foreground=FG_TEXT)

        style.configure("TButton", background=ACCENT, foreground="#ffffff",
                         font=(FONT_FAMILY, 10, "bold"), padding=(13, 7), borderwidth=0)
        style.map("TButton", background=[("active", "#6a4bee")])
        style.configure("Danger.TButton", background=RED, foreground="#ffffff")
        style.map("Danger.TButton", background=[("active", "#e05555")])
        style.configure("Success.TButton", background=GREEN, foreground="#08120c")
        style.map("Success.TButton", background=[("active", "#33c084")])

        style.configure("TCombobox", fieldbackground=BG_CARD_ALT, background=BG_CARD_ALT,
                         foreground=FG_TEXT, arrowcolor=FG_TEXT)
        style.configure("TEntry", fieldbackground=BG_CARD_ALT, foreground=FG_TEXT,
                         insertcolor=FG_TEXT)
        style.configure("TRadiobutton", background=BG_MAIN, foreground=FG_TEXT,
                         font=(FONT_FAMILY, 10))
        style.map("TRadiobutton", background=[("active", BG_MAIN)])

        style.configure("Treeview", background=BG_CARD, fieldbackground=BG_CARD,
                         foreground=FG_TEXT, rowheight=28, borderwidth=0,
                         font=(FONT_FAMILY, 10))
        style.configure("Treeview.Heading", background=BG_CARD_ALT, foreground=FG_TEXT,
                         font=(FONT_FAMILY, 10, "bold"), relief="flat")
        style.map("Treeview", background=[("selected", ACCENT)])

        style.configure("TProgressbar", troughcolor=BG_CARD_ALT, background=ACCENT2,
                         thickness=14)

    def _build_layout(self):
        header = tk.Frame(self, bg=BG_MAIN)
        header.pack(fill="x", padx=24, pady=(20, 10))
        tk.Label(header, text="💰 مدیریت خرج و مخارج", bg=BG_MAIN, fg=FG_TEXT,
                  font=(FONT_FAMILY, 20, "bold")).pack(side="right")
        tk.Label(header, text="درآمد، هزینه، بودجه و گزارش — تقویم شمسی",
                  bg=BG_MAIN, fg=FG_MUTED, font=(FONT_FAMILY, 11)).pack(side="right", padx=12)

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=20, pady=10)
        self.nb = nb

        self.tab_dashboard = ttk.Frame(nb)
        self.tab_add = ttk.Frame(nb)
        self.tab_history = ttk.Frame(nb)
        self.tab_budget = ttk.Frame(nb)
        self.tab_report = ttk.Frame(nb)
        self.tab_settings = ttk.Frame(nb)

        nb.add(self.tab_dashboard, text="📊 داشبورد")
        nb.add(self.tab_add, text="➕ ثبت تراکنش")
        nb.add(self.tab_history, text="📜 تاریخچه")
        nb.add(self.tab_budget, text="🎯 بودجه‌بندی")
        nb.add(self.tab_report, text="📈 گزارش")
        nb.add(self.tab_settings, text="⚙️ پشتیبان‌گیری")

        nb.bind("<<NotebookTabChanged>>", lambda e: self.refresh_all())

        self._build_dashboard(self.tab_dashboard)
        self._build_add_form(self.tab_add)
        self._build_history(self.tab_history)
        self._build_budget(self.tab_budget)
        self._build_report(self.tab_report)
        self._build_settings(self.tab_settings)

    def _build_dashboard(self, parent):
        top = tk.Frame(parent, bg=BG_MAIN)
        top.pack(fill="x", pady=(10, 12))
        tk.Label(top, text="ماه جاری:", bg=BG_MAIN, fg=FG_MUTED,
                  font=(FONT_FAMILY, 11)).pack(side="right", padx=(0, 8))
        self.dash_month_lbl = tk.Label(top, text="", bg=BG_MAIN, fg=ACCENT2,
                                        font=(FONT_FAMILY, 13, "bold"))
        self.dash_month_lbl.pack(side="right")

        self.dash_warning_lbl = tk.Label(parent, text="", bg=BG_MAIN, fg=YELLOW,
                                          font=(FONT_FAMILY, 10, "bold"), justify="right",
                                          anchor="e", wraplength=1000)
        self.dash_warning_lbl.pack(fill="x", pady=(0, 10))

        cards = tk.Frame(parent, bg=BG_MAIN)
        cards.pack(fill="x")
        cards.grid_columnconfigure((0, 1, 2), weight=1, uniform="c")

        self.card_income = StatCard(cards, "درآمد این ماه", GREEN)
        self.card_expense = StatCard(cards, "هزینه این ماه", RED)
        self.card_saving = StatCard(cards, "پس‌انداز این ماه", BLUE)
        self.card_income.grid(row=0, column=2, sticky="nsew", padx=6)
        self.card_expense.grid(row=0, column=1, sticky="nsew", padx=6)
        self.card_saving.grid(row=0, column=0, sticky="nsew", padx=6)

        chart_wrap = tk.Frame(parent, bg=BG_CARD)
        chart_wrap.pack(fill="both", expand=True, pady=(20, 0))
        tk.Label(chart_wrap, text="تفکیک هزینه‌ها بر اساس دسته (ماه جاری)",
                  bg=BG_CARD, fg=FG_TEXT, font=(FONT_FAMILY, 12, "bold")).pack(
                      anchor="e", padx=16, pady=(12, 0))
        self.dash_chart_holder = tk.Frame(chart_wrap, bg=BG_CARD)
        self.dash_chart_holder.pack(fill="both", expand=True, padx=10, pady=10)

    def refresh_dashboard(self):
        today = jtoday()
        rows = self.db.for_month(today.year, today.month)
        income = sum(r[2] for r in rows if r[1] == "income")
        expense = sum(r[2] for r in rows if r[1] == "expense")
        saving = income - expense

        self.dash_month_lbl.config(text=month_label(today.year, today.month))
        self.card_income.set_value(fmt_money(income))
        self.card_expense.set_value(fmt_money(expense))
        self.card_saving.set_value(fmt_money(saving))
        self.card_saving.value_lbl.config(fg=GREEN if saving >= 0 else RED)

        cat_totals = defaultdict(float)
        for r in rows:
            if r[1] == "expense":
                cat_totals[r[3]] += r[2]

        budgets = self.db.get_budgets()
        over = []
        for cat, spent in cat_totals.items():
            b = budgets.get(cat, 0)
            if b > 0 and spent > b:
                pct = spent / b * 100
                over.append(f"{cat} ({pct:.0f}٪ از {fmt_money(b)})")
        if over:
            self.dash_warning_lbl.config(
                text="⚠️ عبور از بودجه در این دسته‌ها:  " + "   •   ".join(over))
        else:
            self.dash_warning_lbl.config(text="")

        for w in self.dash_chart_holder.winfo_children():
            w.destroy()

        if not HAS_MPL:
            tk.Label(self.dash_chart_holder, text="برای نمایش نمودار: pip install matplotlib",
                      bg=BG_CARD, fg=FG_MUTED, font=(FONT_FAMILY, 10)).pack(expand=True)
            return
        if not cat_totals:
            tk.Label(self.dash_chart_holder, text="در این ماه هنوز هزینه‌ای ثبت نشده است.",
                      bg=BG_CARD, fg=FG_MUTED, font=(FONT_FAMILY, 11)).pack(expand=True)
            return

        fig = self._make_pie_figure(cat_totals, figsize=(5.5, 3.6))
        canvas = FigureCanvasTkAgg(fig, master=self.dash_chart_holder)
        canvas.draw()
        canvas.get_tk_widget().pack(fill="both", expand=True)

    def _make_pie_figure(self, cat_totals, figsize=(5.5, 3.6)):
        fig = Figure(figsize=figsize, dpi=100)
        fig.patch.set_facecolor(BG_CARD)
        ax = fig.add_subplot(111)
        ax.set_facecolor(BG_CARD)
        
        labels = list(cat_totals.keys())
        values = list(cat_totals.values())
        colors = [ACCENT, ACCENT2, GREEN, RED, BLUE, YELLOW, "#ff9f6b", "#c792ea", "#82aaff"]
        
        wedges, _texts, autotexts = ax.pie(
            values, autopct="%1.0f%%", colors=colors,
            textprops={"color": "#ffffff", "fontsize": 9},
            startangle=90, wedgeprops={"linewidth": 1, "edgecolor": BG_CARD},
        )
        
        # تنظیم فونت برای برچسب‌های درصد
        if HAS_MPL:
            font_prop = fm.FontProperties(family=PERSIAN_FONT)
            for autotext in autotexts:
                autotext.set_fontproperties(font_prop)
        
        ax.legend(wedges, labels, loc="center left", bbox_to_anchor=(1.0, 0.5),
                   frameon=False, labelcolor=FG_TEXT, fontsize=8)
        fig.tight_layout()
        return fig

    def _build_add_form(self, parent):
        wrap = tk.Frame(parent, bg=BG_CARD)
        wrap.pack(fill="x", padx=40, pady=40)

        self.kind_var = tk.StringVar(value="expense")
        kind_frame = tk.Frame(wrap, bg=BG_CARD)
        kind_frame.pack(fill="x", padx=20, pady=(20, 10))
        ttk.Radiobutton(kind_frame, text="هزینه", variable=self.kind_var,
                        value="expense", command=self._refresh_category_list).pack(side="right", padx=10)
        ttk.Radiobutton(kind_frame, text="درآمد", variable=self.kind_var,
                        value="income", command=self._refresh_category_list).pack(side="right", padx=10)

        grid = tk.Frame(wrap, bg=BG_CARD)
        grid.pack(fill="x", padx=20, pady=10)

        def row(label, widget_factory, r):
            tk.Label(grid, text=label, bg=BG_CARD, fg=FG_TEXT,
                      font=(FONT_FAMILY, 11)).grid(row=r, column=1, sticky="e", padx=10, pady=10)
            w = widget_factory()
            w.grid(row=r, column=0, sticky="ew", padx=10, pady=10)
            return w

        grid.grid_columnconfigure(0, weight=1)

        self.amount_var = tk.StringVar()
        row("مبلغ (تومان)", lambda: ttk.Entry(
            grid, textvariable=self.amount_var, font=(FONT_FAMILY, 11)), 0)

        self.category_var = tk.StringVar()
        self.category_combo = row("دسته‌بندی", lambda: ttk.Combobox(
            grid, textvariable=self.category_var, state="readonly",
            values=EXPENSE_CATEGORIES, font=(FONT_FAMILY, 11)), 1)

        self.note_var = tk.StringVar()
        row("توضیحات (اختیاری)", lambda: ttk.Entry(
            grid, textvariable=self.note_var, font=(FONT_FAMILY, 11)), 2)

        self.date_var = tk.StringVar(value=jalali_str(jtoday()))
        row("تاریخ شمسی (YYYY-MM-DD)", lambda: ttk.Entry(
            grid, textvariable=self.date_var, font=(FONT_FAMILY, 11)), 3)
        tk.Label(grid, text="مثال: 1405-03-11", bg=BG_CARD, fg=FG_MUTED,
                  font=(FONT_FAMILY, 9)).grid(row=4, column=0, sticky="w", padx=10)

        self._refresh_category_list()

        btn_frame = tk.Frame(wrap, bg=BG_CARD)
        btn_frame.pack(fill="x", padx=20, pady=(10, 24))
        ttk.Button(btn_frame, text="ثبت تراکنش", command=self._submit_transaction).pack(side="right")
        ttk.Button(btn_frame, text="امروز", command=self._set_today).pack(side="right", padx=8)
        self.add_status_lbl = tk.Label(btn_frame, text="", bg=BG_CARD, fg=GREEN,
                                        font=(FONT_FAMILY, 10))
        self.add_status_lbl.pack(side="right", padx=14)

    def _set_today(self):
        self.date_var.set(jalali_str(jtoday()))

    def _refresh_category_list(self):
        cats = EXPENSE_CATEGORIES if self.kind_var.get() == "expense" else INCOME_CATEGORIES
        self.category_combo["values"] = cats
        if self.category_var.get() not in cats:
            self.category_var.set(cats[0])

    def _validate_transaction_form(self, amount_raw, date_raw, category):
        try:
            amount = float(amount_raw.strip().replace(",", ""))
            if amount <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("خطا", "مبلغ را به‌صورت یک عدد مثبت وارد کنید.")
            return None
        if not HAS_JDATETIME:
            messagebox.showerror("خطا", "کتابخانه jdatetime نصب نیست:\npip install jdatetime")
            return None
        try:
            jd = parse_jalali(date_raw)
        except ValueError:
            messagebox.showerror("خطا", "تاریخ شمسی نامعتبر است. فرمت درست: 1405-03-11")
            return None
        if not category:
            messagebox.showerror("خطا", "یک دسته‌بندی انتخاب کنید.")
            return None
        return amount, jalali_str(jd)

    def _submit_transaction(self):
        result = self._validate_transaction_form(
            self.amount_var.get(), self.date_var.get(), self.category_var.get())
        if result is None:
            return
        amount, tdate = result
        self.db.add(self.kind_var.get(), amount, self.category_var.get(),
                    self.note_var.get().strip(), tdate)
        self.amount_var.set("")
        self.note_var.set("")
        self.add_status_lbl.config(text="✓ با موفقیت ثبت شد")
        self.after(2000, lambda: self.add_status_lbl.config(text=""))
        self.refresh_all()

    def _build_history(self, parent):
        filt = tk.Frame(parent, bg=BG_MAIN)
        filt.pack(fill="x", pady=(10, 10))

        tk.Label(filt, text="فیلتر بر اساس ماه:", bg=BG_MAIN, fg=FG_MUTED,
                  font=(FONT_FAMILY, 10)).pack(side="right", padx=6)
        self.hist_year_cb = ttk.Combobox(filt, state="readonly", width=8, font=(FONT_FAMILY, 10))
        self.hist_year_cb.pack(side="right", padx=4)
        self.hist_month_cb = ttk.Combobox(filt, state="readonly", width=10, font=(FONT_FAMILY, 10),
                                           values=JALALI_MONTH_NAMES)
        self.hist_month_cb.pack(side="right", padx=4)
        ttk.Button(filt, text="اعمال فیلتر", command=self.refresh_history).pack(side="right", padx=8)
        ttk.Button(filt, text="نمایش همه", command=lambda: self.refresh_history(show_all=True)).pack(
            side="right", padx=4)
        ttk.Button(filt, text="📥 خروجی اکسل", command=self._export_history_excel).pack(
            side="left", padx=4)

        cols = ("date", "kind", "category", "amount", "note", "id")
        self.tree = ttk.Treeview(parent, columns=cols, show="headings", selectmode="browse")
        headings = {"date": "تاریخ شمسی", "kind": "نوع", "category": "دسته‌بندی",
                    "amount": "مبلغ", "note": "توضیحات", "id": "شناسه"}
        widths = {"date": 110, "kind": 70, "category": 140, "amount": 110, "note": 230, "id": 0}
        for c in cols:
            self.tree.heading(c, text=headings[c])
            self.tree.column(c, width=widths[c], anchor="center", stretch=(c != "id"))
        self.tree.column("id", width=0, stretch=False)
        self.tree["displaycolumns"] = ("date", "kind", "category", "amount", "note")
        self.tree.pack(fill="both", expand=True, pady=6)
        self.tree.bind("<Double-1>", lambda e: self._edit_selected())

        self.tree.tag_configure("income", foreground=GREEN)
        self.tree.tag_configure("expense", foreground=RED)

        bottom = tk.Frame(parent, bg=BG_MAIN)
        bottom.pack(fill="x", pady=6)
        ttk.Button(bottom, text="حذف تراکنش", style="Danger.TButton",
                   command=self._delete_selected).pack(side="right")
        ttk.Button(bottom, text="✏️ ویرایش تراکنش", command=self._edit_selected).pack(
            side="right", padx=8)
        self.hist_summary_lbl = tk.Label(bottom, text="", bg=BG_MAIN, fg=FG_MUTED,
                                          font=(FONT_FAMILY, 10))
        self.hist_summary_lbl.pack(side="left")

        self._current_history_rows = []

    def refresh_history(self, show_all=False):
        years = self.db.years_present()
        self.hist_year_cb["values"] = years
        today = jtoday()
        if not self.hist_year_cb.get():
            self.hist_year_cb.set(today.year)
        if not self.hist_month_cb.get():
            self.hist_month_cb.set(JALALI_MONTH_NAMES[today.month - 1])

        for i in self.tree.get_children():
            self.tree.delete(i)

        if show_all:
            rows = self.db.all()
        else:
            try:
                yr = int(self.hist_year_cb.get())
                mo = JALALI_MONTH_NAMES.index(self.hist_month_cb.get()) + 1
            except (ValueError, IndexError):
                yr, mo = today.year, today.month
            rows = self.db.for_month(yr, mo)

        self._current_history_rows = rows
        income = expense = 0
        for tid, kind, amount, category, note, tdate in rows:
            kind_fa = "درآمد" if kind == "income" else "هزینه"
            self.tree.insert("", "end", values=(tdate, kind_fa, category,
                                                  fmt_money(amount), note or "-", tid),
                              tags=(kind,))
            if kind == "income":
                income += amount
            else:
                expense += amount

        self.hist_summary_lbl.config(
            text=f"جمع درآمد: {fmt_money(income)}   |   جمع هزینه: {fmt_money(expense)}   |   "
                 f"پس‌انداز: {fmt_money(income - expense)}   |   تعداد تراکنش: {len(rows)}"
        )

    def _get_selected_id(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("توجه", "ابتدا یک تراکنش را از لیست انتخاب کنید.")
            return None
        return self.tree.item(sel[0])["values"][-1]

    def _delete_selected(self):
        tid = self._get_selected_id()
        if tid is None:
            return
        if messagebox.askyesno("تأیید حذف", "از حذف این تراکنش مطمئن هستید؟"):
            self.db.delete(tid)
            self.refresh_all()

    def _edit_selected(self):
        tid = self._get_selected_id()
        if tid is None:
            return
        row = self.db.get(tid)
        if not row:
            return
        self._open_edit_dialog(row)

    def _open_edit_dialog(self, row):
        tid, kind, amount, category, note, tdate = row
        if self.edit_dialog and self.edit_dialog.winfo_exists():
            self.edit_dialog.destroy()

        dlg = tk.Toplevel(self)
        self.edit_dialog = dlg
        dlg.title("ویرایش تراکنش")
        dlg.configure(bg=BG_CARD)
        dlg.geometry("420x380")
        dlg.transient(self)
        dlg.grab_set()

        kind_var = tk.StringVar(value=kind)
        amount_var = tk.StringVar(value=fmt_money(amount))
        note_var = tk.StringVar(value=note or "")
        date_var = tk.StringVar(value=tdate)
        category_var = tk.StringVar(value=category)

        kf = tk.Frame(dlg, bg=BG_CARD)
        kf.pack(fill="x", padx=20, pady=(20, 10))

        def cats_for_kind():
            return EXPENSE_CATEGORIES if kind_var.get() == "expense" else INCOME_CATEGORIES

        def refresh_cats():
            cat_combo["values"] = cats_for_kind()
            if category_var.get() not in cat_combo["values"]:
                category_var.set(cat_combo["values"][0])

        ttk.Radiobutton(kf, text="هزینه", variable=kind_var, value="expense",
                        command=refresh_cats).pack(side="right", padx=10)
        ttk.Radiobutton(kf, text="درآمد", variable=kind_var, value="income",
                        command=refresh_cats).pack(side="right", padx=10)

        grid = tk.Frame(dlg, bg=BG_CARD)
        grid.pack(fill="x", padx=20)
        grid.grid_columnconfigure(0, weight=1)

        tk.Label(grid, text="مبلغ", bg=BG_CARD, fg=FG_TEXT).grid(row=0, column=1, sticky="e", pady=8)
        ttk.Entry(grid, textvariable=amount_var).grid(row=0, column=0, sticky="ew", pady=8, padx=8)

        tk.Label(grid, text="دسته‌بندی", bg=BG_CARD, fg=FG_TEXT).grid(row=1, column=1, sticky="e", pady=8)
        cat_combo = ttk.Combobox(grid, textvariable=category_var, state="readonly",
                                  values=cats_for_kind())
        cat_combo.grid(row=1, column=0, sticky="ew", pady=8, padx=8)

        tk.Label(grid, text="توضیحات", bg=BG_CARD, fg=FG_TEXT).grid(row=2, column=1, sticky="e", pady=8)
        ttk.Entry(grid, textvariable=note_var).grid(row=2, column=0, sticky="ew", pady=8, padx=8)

        tk.Label(grid, text="تاریخ (YYYY-MM-DD)", bg=BG_CARD, fg=FG_TEXT).grid(
            row=3, column=1, sticky="e", pady=8)
        ttk.Entry(grid, textvariable=date_var).grid(row=3, column=0, sticky="ew", pady=8, padx=8)

        def save():
            result = self._validate_transaction_form(
                amount_var.get(), date_var.get(), category_var.get())
            if result is None:
                return
            amt, tdate2 = result
            self.db.update(tid, kind_var.get(), amt, category_var.get(),
                           note_var.get().strip(), tdate2)
            dlg.destroy()
            self.refresh_all()

        btns = tk.Frame(dlg, bg=BG_CARD)
        btns.pack(fill="x", padx=20, pady=20)
        ttk.Button(btns, text="ذخیره تغییرات", command=save).pack(side="right")
        ttk.Button(btns, text="انصراف", command=dlg.destroy).pack(side="right", padx=8)

    def _export_history_excel(self):
        rows = self._current_history_rows
        if not rows:
            messagebox.showinfo("توجه", "چیزی برای خروجی گرفتن وجود ندارد.")
            return
        self._export_transactions_excel(rows, default_name="تاریخچه")

    def _build_budget(self, parent):
        top = tk.Frame(parent, bg=BG_MAIN)
        top.pack(fill="x", pady=(10, 6))
        tk.Label(top, text="بودجه‌ی ماهانه برای هر دسته‌ی هزینه را تعیین کنید",
                  bg=BG_MAIN, fg=FG_MUTED, font=(FONT_FAMILY, 11)).pack(side="right")

        card = tk.Frame(parent, bg=BG_CARD)
        card.pack(fill="both", expand=True, pady=8)

        head = tk.Frame(card, bg=BG_CARD)
        head.pack(fill="x", padx=16, pady=(14, 4))
        for text, col in [("دسته‌بندی", 4), ("وضعیت", 1), ("پیشرفت", 2), ("بودجه (تومان)", 3)]:
            tk.Label(head, text=text, bg=BG_CARD, fg=FG_MUTED,
                      font=(FONT_FAMILY, 9, "bold")).grid(row=0, column=col, padx=8)
        head.grid_columnconfigure(0, weight=1)

        rows_frame = tk.Frame(card, bg=BG_CARD)
        rows_frame.pack(fill="x", padx=16)

        for cat in EXPENSE_CATEGORIES:
            r = BudgetRow(rows_frame, cat, on_change=self._save_budgets)
            r.pack(fill="x", pady=4)
            self.budget_rows[cat] = r

        btn_frame = tk.Frame(card, bg=BG_CARD)
        btn_frame.pack(fill="x", padx=16, pady=16)
        ttk.Button(btn_frame, text="ذخیره‌ی همه‌ی بودجه‌ها", command=self._save_budgets).pack(
            side="right")
        self.budget_status_lbl = tk.Label(btn_frame, text="", bg=BG_CARD, fg=GREEN,
                                           font=(FONT_FAMILY, 10))
        self.budget_status_lbl.pack(side="right", padx=14)

    def _save_budgets(self):
        for cat, row_widget in self.budget_rows.items():
            self.db.set_budget(cat, row_widget.get_budget())
        self.budget_status_lbl.config(text="✓ ذخیره شد")
        self.after(1500, lambda: self.budget_status_lbl.config(text=""))
        self.refresh_budget()
        self.refresh_dashboard()

    def refresh_budget(self):
        budgets = self.db.get_budgets()
        today = jtoday()
        rows = self.db.for_month(today.year, today.month)
        spent_by_cat = defaultdict(float)
        for r in rows:
            if r[1] == "expense":
                spent_by_cat[r[3]] += r[2]

        for cat, row_widget in self.budget_rows.items():
            b = budgets.get(cat, 0)
            row_widget.set_budget(b)
            row_widget.set_spent(spent_by_cat.get(cat, 0), b)

    def _build_report(self, parent):
        top = tk.Frame(parent, bg=BG_MAIN)
        top.pack(fill="x", pady=(10, 10))
        tk.Label(top, text="سال شمسی:", bg=BG_MAIN, fg=FG_MUTED,
                  font=(FONT_FAMILY, 10)).pack(side="right", padx=6)
        self.report_year_cb = ttk.Combobox(top, state="readonly", width=8, font=(FONT_FAMILY, 10))
        self.report_year_cb.pack(side="right")
        ttk.Button(top, text="بروزرسانی نمودار", command=self.refresh_report).pack(side="right", padx=10)
        ttk.Button(top, text="📥 خروجی اکسل", command=self._export_report_excel).pack(
            side="left", padx=4)
        ttk.Button(top, text="📄 خروجی PDF", command=self._export_report_pdf).pack(
            side="left", padx=4)

        self.report_chart_holder = tk.Frame(parent, bg=BG_CARD)
        self.report_chart_holder.pack(fill="both", expand=True, pady=10)

    def _compute_year_data(self, year):
        incomes, expenses, savings = [], [], []
        for m in range(1, 13):
            rows = self.db.for_month(year, m)
            inc = sum(r[2] for r in rows if r[1] == "income")
            exp = sum(r[2] for r in rows if r[1] == "expense")
            incomes.append(inc)
            expenses.append(exp)
            savings.append(inc - exp)
        return incomes, expenses, savings

    def refresh_report(self):
        years = self.db.years_present()
        self.report_year_cb["values"] = years
        if not self.report_year_cb.get():
            self.report_year_cb.set(jtoday().year)
        try:
            year = int(self.report_year_cb.get())
        except ValueError:
            year = jtoday().year

        for w in self.report_chart_holder.winfo_children():
            w.destroy()

        if not HAS_MPL:
            tk.Label(self.report_chart_holder, text="برای نمایش نمودار: pip install matplotlib",
                      bg=BG_CARD, fg=FG_MUTED, font=(FONT_FAMILY, 10)).pack(expand=True)
            return

        incomes, expenses, savings = self._compute_year_data(year)
        fig = self._make_year_bar_figure(year, incomes, expenses, savings)
        canvas = FigureCanvasTkAgg(fig, master=self.report_chart_holder)
        canvas.draw()
        canvas.get_tk_widget().pack(fill="both", expand=True)

        total_income, total_expense = sum(incomes), sum(expenses)
        tk.Label(
            self.report_chart_holder,
            text=(f"جمع کل سال {year}   —   درآمد: {fmt_money(total_income)}   |   "
                  f"هزینه: {fmt_money(total_expense)}   |   "
                  f"پس‌انداز خالص: {fmt_money(total_income - total_expense)}"),
            bg=BG_CARD, fg=FG_TEXT, font=(FONT_FAMILY, 11, "bold"),
        ).pack(pady=(0, 12))

    def _make_year_bar_figure(self, year, incomes, expenses, savings):
        fig = Figure(figsize=(9, 4.2), dpi=100)
        fig.patch.set_facecolor(BG_CARD)
        ax = fig.add_subplot(111)
        ax.set_facecolor(BG_CARD)
        x = range(12)
        width = 0.27
        
        # استفاده از فونت فارسی برای برچسب‌ها
        if HAS_MPL:
            font_prop = fm.FontProperties(family=PERSIAN_FONT)
        
        ax.bar([i - width for i in x], incomes, width=width, label="درآمد", color=GREEN)
        ax.bar(list(x), expenses, width=width, label="هزینه", color=RED)
        ax.bar([i + width for i in x], savings, width=width, label="پس‌انداز", color=BLUE)
        ax.set_xticks(list(x))
        
        # استفاده از نام کامل ماه‌ها به جای نام‌های کوتاه
        if HAS_MPL:
            ax.set_xticklabels(JALALI_MONTH_NAMES, color=FG_TEXT, fontsize=9, fontproperties=font_prop)
        else:
            ax.set_xticklabels(JALALI_MONTH_NAMES, color=FG_TEXT, fontsize=9)
        
        ax.tick_params(colors=FG_TEXT)
        
        # تنظیم فرمت اعداد برای محور y
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: format(int(x), ',')))
        
        for spine in ax.spines.values():
            spine.set_color(BG_CARD_ALT)
        ax.axhline(0, color=BG_CARD_ALT, linewidth=1)
        
        # تنظیم فونت برای افسانه و عنوان
        if HAS_MPL:
            ax.legend(facecolor=BG_CARD, edgecolor=BG_CARD, labelcolor=FG_TEXT, prop=font_prop)
            ax.set_title(f"گزارش سالانه {year} (تقویم شمسی)", color=FG_TEXT, fontsize=12, fontproperties=font_prop)
        else:
            ax.legend(facecolor=BG_CARD, edgecolor=BG_CARD, labelcolor=FG_TEXT)
            ax.set_title(f"گزارش سالانه {year} (تقویم شمسی)", color=FG_TEXT, fontsize=12)
        
        fig.tight_layout()
        return fig

    def _export_transactions_excel(self, rows, default_name):
        if not HAS_OPENPYXL:
            messagebox.showerror("خطا", "کتابخانه openpyxl نصب نیست:\npip install openpyxl")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile=f"{default_name}.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "تراکنش‌ها"
        ws.sheet_view.rightToLeft = True

        headers = ["تاریخ شمسی", "نوع", "دسته‌بندی", "مبلغ", "توضیحات"]
        ws.append(headers)
        header_fill = PatternFill(start_color="7C5CFF", end_color="7C5CFF", fill_type="solid")
        for cell in ws[1]:
            cell.font = XLFont(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for tid, kind, amount, category, note, tdate in rows:
            ws.append([tdate, "درآمد" if kind == "income" else "هزینه",
                       category, amount, note or ""])

        for i, w in enumerate([14, 10, 20, 14, 30], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

        try:
            wb.save(path)
            messagebox.showinfo("موفق", f"فایل اکسل ذخیره شد:\n{path}")
        except Exception as e:
            messagebox.showerror("خطا", f"ذخیره‌سازی ناموفق بود:\n{e}")

    def _export_report_excel(self):
        if not HAS_OPENPYXL:
            messagebox.showerror("خطا", "کتابخانه openpyxl نصب نیست:\npip install openpyxl")
            return
        try:
            year = int(self.report_year_cb.get())
        except ValueError:
            year = jtoday().year
        incomes, expenses, savings = self._compute_year_data(year)

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile=f"گزارش-{year}.xlsx",
            filetypes=[("Excel", "*.xlsx")])
        if not path:
            return

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "خلاصه ماهانه"
        ws1.sheet_view.rightToLeft = True
        ws1.append(["ماه", "درآمد", "هزینه", "پس‌انداز"])
        header_fill = PatternFill(start_color="7C5CFF", end_color="7C5CFF", fill_type="solid")
        for cell in ws1[1]:
            cell.font = XLFont(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for i in range(12):
            ws1.append([JALALI_MONTH_NAMES[i], incomes[i], expenses[i], savings[i]])
        ws1.append(["جمع کل", sum(incomes), sum(expenses), sum(savings)])
        for cell in ws1[ws1.max_row]:
            cell.font = XLFont(bold=True)
        for i, w in enumerate([16, 14, 14, 14], start=1):
            ws1.column_dimensions[get_column_letter(i)].width = w

        ws2 = wb.create_sheet("تراکنش‌های سال")
        ws2.sheet_view.rightToLeft = True
        ws2.append(["تاریخ شمسی", "نوع", "دسته‌بندی", "مبلغ", "توضیحات"])
        for cell in ws2[1]:
            cell.font = XLFont(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for tid, kind, amount, category, note, tdate in self.db.for_year(year):
            ws2.append([tdate, "درآمد" if kind == "income" else "هزینه",
                        category, amount, note or ""])
        for i, w in enumerate([14, 10, 20, 14, 30], start=1):
            ws2.column_dimensions[get_column_letter(i)].width = w

        try:
            wb.save(path)
            messagebox.showinfo("موفق", f"فایل اکسل ذخیره شد:\n{path}")
        except Exception as e:
            messagebox.showerror("خطا", f"ذخیره‌سازی ناموفق بود:\n{e}")

    def _export_report_pdf(self):
        if not HAS_REPORTLAB:
            messagebox.showerror("خطا", "کتابخانه reportlab نصب نیست:\npip install reportlab")
            return
        try:
            year = int(self.report_year_cb.get())
        except ValueError:
            year = jtoday().year

        path = filedialog.asksaveasfilename(
            defaultextension=".pdf", initialfile=f"گزارش-{year}.pdf",
            filetypes=[("PDF", "*.pdf")])
        if not path:
            return

        font_name = register_persian_font()
        if font_name == "Helvetica":
            messagebox.showwarning(
                "فونت فارسی یافت نشد",
                "فونت فارسی مناسبی پیدا نشد و ممکن است متن فارسی در PDF درست نمایش داده "
                "نشود.\nبرای رفع این مشکل، فونت Vazirmatn-Regular.ttf را دانلود و کنار "
                "فایل expense_manager.py قرار دهید.")

        incomes, expenses, savings = self._compute_year_data(year)

        doc = SimpleDocTemplate(path, pagesize=A4,
                                 topMargin=1.5 * cm, bottomMargin=1.5 * cm)
        story = []

        title_style = ParagraphStyle("title", fontName=font_name, fontSize=18,
                                      alignment=1, spaceAfter=14, textColor=rl_colors.HexColor("#1a1e33"))
        normal_style = ParagraphStyle("normal", fontName=font_name, fontSize=11,
                                       alignment=1, spaceAfter=10)

        story.append(Paragraph(rtl(f"گزارش سالانه {year} — مدیریت خرج و مخارج"), title_style))
        total_income, total_expense = sum(incomes), sum(expenses)
        story.append(Paragraph(
            rtl(f"جمع درآمد: {fmt_money(total_income)}   |   جمع هزینه: {fmt_money(total_expense)}"
                f"   |   پس‌انداز خالص: {fmt_money(total_income - total_expense)}"),
            normal_style))
        story.append(Spacer(1, 10))

        if HAS_MPL:
            fig = self._make_year_bar_figure(year, incomes, expenses, savings)
            buf = BytesIO()
            fig.savefig(buf, format="png", dpi=140, facecolor=fig.get_facecolor())
            buf.seek(0)
            story.append(RLImage(buf, width=16 * cm, height=7.5 * cm))
            story.append(Spacer(1, 14))

        table_data = [[rtl("ماه"), rtl("درآمد"), rtl("هزینه"), rtl("پس‌انداز")]]
        for i in range(12):
            table_data.append([
                rtl(JALALI_MONTH_NAMES[i]), fmt_money(incomes[i]),
                fmt_money(expenses[i]), fmt_money(savings[i]),
            ])
        table_data.append([rtl("جمع کل"), fmt_money(total_income),
                            fmt_money(total_expense), fmt_money(total_income - total_expense)])

        tbl = Table(table_data, colWidths=[4 * cm, 4 * cm, 4 * cm, 4 * cm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#7c5cff")),
            ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.HexColor("#cccccc")),
            ("BACKGROUND", (0, -1), (-1, -1), rl_colors.HexColor("#e8e6ff")),
            ("FONTNAME", (0, -1), (-1, -1), font_name),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2),
             [rl_colors.white, rl_colors.HexColor("#f5f5fb")]),
        ]))
        story.append(tbl)

        try:
            doc.build(story)
            messagebox.showinfo("موفق", f"فایل PDF ذخیره شد:\n{path}")
        except Exception as e:
            messagebox.showerror("خطا", f"ساخت PDF ناموفق بود:\n{e}")

    def _build_settings(self, parent):
        info = tk.Frame(parent, bg=BG_CARD)
        info.pack(fill="x", pady=(10, 12))
        tk.Label(info, text=f"محل پایگاه‌داده: {DB_PATH}", bg=BG_CARD, fg=FG_MUTED,
                  font=(FONT_FAMILY, 9), anchor="e", justify="right").pack(
                      fill="x", padx=16, pady=(12, 2))
        tk.Label(info, text="یک نسخه‌ی پشتیبان به‌صورت خودکار روزی یک‌بار (در اولین اجرای "
                             "برنامه در آن روز) گرفته می‌شود.",
                  bg=BG_CARD, fg=FG_MUTED, font=(FONT_FAMILY, 9), anchor="e",
                  justify="right").pack(fill="x", padx=16, pady=(0, 12))

        btns = tk.Frame(parent, bg=BG_MAIN)
        btns.pack(fill="x", pady=(0, 10))
        ttk.Button(btns, text="📦 پشتیبان‌گیری دستی الان", style="Success.TButton",
                   command=self._manual_backup).pack(side="right")
        ttk.Button(btns, text="♻️ بازیابی از پشتیبان انتخاب‌شده", style="Danger.TButton",
                   command=self._restore_backup).pack(side="right", padx=8)
        ttk.Button(btns, text="بروزرسانی لیست", command=self.refresh_settings).pack(
            side="left")

        cols = ("file", "size")
        self.backup_tree = ttk.Treeview(parent, columns=cols, show="headings",
                                         selectmode="browse", height=12)
        self.backup_tree.heading("file", text="نام فایل پشتیبان")
        self.backup_tree.heading("size", text="حجم")
        self.backup_tree.column("file", width=420, anchor="center")
        self.backup_tree.column("size", width=100, anchor="center")
        self.backup_tree.pack(fill="both", expand=True, pady=6)

    def _manual_backup(self):
        dest = backup_now()
        if dest:
            messagebox.showinfo("موفق", f"پشتیبان ذخیره شد:\n{dest}")
        else:
            messagebox.showinfo("توجه", "هنوز پایگاه‌داده‌ای برای پشتیبان‌گیری وجود ندارد.")
        self.refresh_settings()

    def _restore_backup(self):
        sel = self.backup_tree.selection()
        if not sel:
            messagebox.showinfo("توجه", "ابتدا یک فایل پشتیبان را از لیست انتخاب کنید.")
            return
        idx = self.backup_tree.index(sel[0])
        backups = list_backups()
        if idx >= len(backups):
            return
        full_path = backups[idx][0]
        if not messagebox.askyesno(
                "تأیید بازیابی",
                "با این کار اطلاعات فعلی جایگزین نسخه‌ی پشتیبان انتخاب‌شده می‌شود.\n"
                "پیشنهاد می‌شود قبل از ادامه یک پشتیبان دستی هم بگیرید.\nادامه می‌دهید؟"):
            return
        try:
            self.db.conn.close()
            shutil.copy2(full_path, DB_PATH)
            self.db.reconnect()
            self.refresh_all()
            messagebox.showinfo("موفق", "بازیابی با موفقیت انجام شد.")
        except Exception as e:
            messagebox.showerror("خطا", f"بازیابی ناموفق بود:\n{e}")

    def refresh_settings(self):
        for i in self.backup_tree.get_children():
            self.backup_tree.delete(i)
        for full_path, name, size in list_backups():
            self.backup_tree.insert("", "end", values=(name, size))

    def refresh_all(self):
        self.refresh_dashboard()
        self.refresh_history()
        self.refresh_budget()
        self.refresh_report()
        self.refresh_settings()


if __name__ == "__main__":
    app = ExpenseApp()
    app.mainloop()